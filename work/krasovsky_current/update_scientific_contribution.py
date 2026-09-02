from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph


SOURCE = Path("/Users/sergeidolin/Downloads/06_Характеристика_соискателя.docx")
OUTPUT = Path("/Users/sergeidolin/RTKLIB/06_Характеристика_соискателя_структурировано.docx")


BLOCKS = [
    (
        "Научная концепция. ",
        "С.В. Долиным сформулирована концепция коллаборативного высокоточного "
        "позиционирования, объединяющая технологии PPP и RTK, и определена методология "
        "исследования.",
    ),
    (
        "Методы и алгоритмы. ",
        "Автором разработаны алгоритм повторной инициализации расширенного фильтра Калмана, "
        "методы учёта дифференциальных кодовых и фазовых задержек и методика PPP-AR. Для "
        "измерений смартфонов созданы адаптивная модель ошибок, робастный вариационно-"
        "байесовский фильтр и методика учёта фазового центра встроенной антенны.",
    ),
    (
        "Программная реализация. ",
        "С.В. Долин модифицировал вычислительное ядро RTKLIB, реализовал разработанные "
        "алгоритмы на языках C и Python и спроектировал архитектуру онлайн-сервиса "
        "апостериорной обработки ГНСС-наблюдений методом PPP-AR.",
    ),
    (
        "Экспериментальное подтверждение. ",
        "Соискатель организовал натурные и вычислительные эксперименты, включая обработку "
        "наблюдений 289 пунктов сети IGS, сопоставил результаты с координатами ITRF2020 и "
        "решениями зарубежных сервисов, выполнил статистическую оценку точности и доли "
        "фиксированных решений.",
    ),
    (
        "Развитие исследования. ",
        "В 2024 году С.В. Долин защитил кандидатскую диссертацию «Разработка методики "
        "коллаборативного позиционирования объектов по сигналам глобальных навигационных "
        "спутниковых систем»; решение диссертационного совета принято единогласно — 14 "
        "голосов «за». Последующие результаты довели предложенную методику до "
        "функционирующей программно-технологической системы.",
    ),
]


def template_rpr(paragraph):
    for run in paragraph.runs:
        if run._r.rPr is not None:
            return deepcopy(run._r.rPr)
    return None


def clear_text_content(paragraph) -> None:
    for child in list(paragraph._p):
        if child.tag in {qn("w:r"), qn("w:hyperlink")}:
            paragraph._p.remove(child)


def add_formatted_run(paragraph, text: str, rpr, bold: bool = False) -> None:
    run = paragraph.add_run(text)
    if rpr is not None:
        if run._r.rPr is not None:
            run._r.remove(run._r.rPr)
        run._r.insert(0, deepcopy(rpr))
    run.bold = bold


def set_labeled_text(paragraph, label: str, body: str) -> None:
    rpr = template_rpr(paragraph)
    clear_text_content(paragraph)
    add_formatted_run(paragraph, label, rpr, bold=True)
    add_formatted_run(paragraph, body, rpr, bold=False)


def main() -> None:
    document = Document(SOURCE)
    heading_index = next(
        i for i, paragraph in enumerate(document.paragraphs)
        if paragraph.text.strip() == "Личный вклад"
    )
    first = document.paragraphs[heading_index + 1]
    dissertation = document.paragraphs[heading_index + 2]

    set_labeled_text(first, *BLOCKS[0])

    for label, body in reversed(BLOCKS[1:4]):
        cloned_xml = deepcopy(first._p)
        first._p.addnext(cloned_xml)
        inserted = Paragraph(cloned_xml, first._parent)
        set_labeled_text(inserted, label, body)

    set_labeled_text(dissertation, *BLOCKS[4])
    document.save(OUTPUT)


if __name__ == "__main__":
    main()
