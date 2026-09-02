from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn


SOURCE = Path("/Users/sergeidolin/RTKLIB/06_Характеристика_соискателя_с_новыми_результатами.docx")
OUTPUT = Path("/Users/sergeidolin/RTKLIB/06_Характеристика_соискателя_итоговая.docx")


MAIN_RESULT = (
    "Предложен механизм повторной инициализации расширенного фильтра Калмана, при котором "
    "координаты относительного решения вводятся в PPP-AR как псевдонаблюдение. При этом "
    "сохраняются оценки фазовых неоднозначностей и атмосферных задержек, что обеспечивает "
    "непрерывность решения при появлении и исчезновении временной базовой станции. Для "
    "двухчастотных смартфонов модифицированы PPP-модели L1/L5 и E1/E5a, разработан "
    "робастный вариационно-байесовский фильтр и методика определения среднего фазового "
    "центра встроенной антенны."
)

METHODS = (
    "Методы и алгоритмы. ",
    "Автор разработал алгоритм смены ролей и повторную инициализацию PPP-AR координатами "
    "относительного решения. Для смартфонов модифицированы модели PPP и разработаны "
    "робастная проверка по критерию Стьюдента, вариационно-байесовская адаптация фильтра и "
    "определение фазового центра антенны.",
)

MOBILE = (
    "Мобильная аппаратура. ",
    "Для Huawei P40 Pro определено смещение среднего фазового центра: 2,7 см влево, 1,3 см "
    "в глубину корпуса и 5,8 см вниз. По десяти двухчасовым сеансам его учёт снизил СКО "
    "высоты: в PPP — с 0,100 до 0,067 м (GPS) и с 0,214 до 0,099 м (GPS + Galileo), в "
    "относительном методе — с 0,146 до 0,092 м и с 0,126 до 0,066 м соответственно.",
)

RECOGNITION = (
    "Результаты представлены на российских и международных конференциях и опубликованы в "
    "рецензируемых российских и зарубежных изданиях; в 2025 году — в журнале Gyroscopy and "
    "Navigation (учёт фазового центра смартфона). С.В. Долин участвовал в "
    "рабочей группе FIG 5.6 «Экономически эффективное позиционирование» в 2019-2021 годах и "
    "входит в состав Российско-китайского Комитета проектов по важному стратегическому "
    "сотрудничеству в области спутниковой навигации."
)


def copy_rpr(paragraph):
    for run in paragraph.runs:
        if run._r.rPr is not None:
            return deepcopy(run._r.rPr)
    return None


def clear_text(paragraph):
    for child in list(paragraph._p):
        if child.tag in {qn("w:r"), qn("w:hyperlink")}:
            paragraph._p.remove(child)


def add_run(paragraph, text, rpr, bold=None):
    run = paragraph.add_run(text)
    if rpr is not None:
        if run._r.rPr is not None:
            run._r.remove(run._r.rPr)
        run._r.insert(0, deepcopy(rpr))
    if bold is not None:
        run.bold = bold


def set_plain(paragraph, text):
    rpr = copy_rpr(paragraph)
    clear_text(paragraph)
    add_run(paragraph, text, rpr)


def set_labeled(paragraph, label, body):
    rpr = copy_rpr(paragraph)
    clear_text(paragraph)
    add_run(paragraph, label, rpr, bold=True)
    add_run(paragraph, body, rpr, bold=False)
    paragraph.paragraph_format.keep_together = True


def after_heading(document, heading, offset=1):
    index = next(i for i, p in enumerate(document.paragraphs) if p.text.strip() == heading)
    return document.paragraphs[index + offset]


def main():
    document = Document(SOURCE)

    set_plain(after_heading(document, "Основные научные результаты", 2), MAIN_RESULT)

    methods = next(p for p in document.paragraphs if p.text.startswith("Методы и алгоритмы."))
    set_labeled(methods, *METHODS)

    mobile = next(p for p in document.paragraphs if p.text.startswith("Мобильная аппаратура."))
    set_labeled(mobile, *MOBILE)

    recognition = after_heading(document, "Признание и образовательная деятельность", 1)
    set_plain(recognition, RECOGNITION)

    document.save(OUTPUT)


if __name__ == "__main__":
    main()
